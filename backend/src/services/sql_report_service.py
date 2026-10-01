"""Báo cáo theo schema SQL 6 bảng, không sử dụng ORM UUID."""
import os
from pathlib import Path
from uuid import uuid4
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.concurrency import run_in_threadpool

EMOTIONS = ('angry', 'disgust', 'fear', 'happy', 'sad', 'surprise', 'neutral')


def safe(value):
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def build_excel(session, metrics, details):
    wb = Workbook()
    info = wb.active
    info.title = 'Thong tin phien'
    info.append(['Thuộc tính', 'Giá trị'])
    for key, value in session.items():
        info.append([key, safe(value)])
    info.append(['Ghi chú', 'Nếu dùng SQL seed thì kết quả là giả lập.'])
    summary = wb.create_sheet('Tong ket')
    summary.append(['Chỉ tiêu', 'Giá trị'])
    if metrics:
        for key, value in metrics.items():
            summary.append([key, safe(value)])
        summary.append(['Lưu ý chỉ số', 'Stress/confidence lấy nguyên từ DB theo thuật toán của nhóm.'])
    else:
        summary.append(['Thông báo', 'Chưa có session_metrics; không tự tạo chỉ số thay thế.'])
    summary.append(['Tổng lượt mặt trong dữ liệu chi tiết', len({r['face_id'] for r in details if r['face_id'] is not None})])
    summary.append(['Cách đếm', 'Lượt phát hiện khuôn mặt, không phải số người khác nhau.'])

    sheet = wb.create_sheet('Chi tiet')
    keys = ['frame_index', 'timestamp_sec', 'frame_path', 'face_id',
            'bbox_x', 'bbox_y', 'bbox_width', 'bbox_height', 'face_confidence',
            'face_image_path', 'emotion', 'score', 'is_dominant']
    sheet.append(keys)
    for row in details:
        sheet.append([safe(row[key]) for key in keys])
    # LEFT JOIN giữ cả frame không có mặt/mặt chưa có điểm cảm xúc.
    note = wb.create_sheet('Huong dan')
    note.append(['Nội dung', 'Giải thích'])
    note.append(['score', 'Giữ nguyên thang điểm DB (0–1 hoặc 0–100); không tự nhân 100.'])
    note.append(['pct_*', 'Tỷ lệ phần trăm đã lưu trong session_metrics.'])
    note.append(['Chi tiết', 'Mỗi dòng là một điểm cảm xúc; frame không có mặt vẫn được liệt kê.'])
    note.append(['Chỉ số', 'Không dùng chỉ số demo để kết luận trạng thái tâm lý.'])
    for ws in wb:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='2563EB')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(65, max(14, max(len(str(c.value or '')) for c in col) + 2))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def save_file(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix('.tmp')
    try:
        temporary.write_bytes(content)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


async def export_sql_excel(db: AsyncSession, session_id: int):
    params = {'session_id': session_id}
    session = (await db.execute(text('SELECT * FROM analysis_sessions WHERE id=:session_id'), params)).mappings().first()
    if session is None:
        raise HTTPException(404, 'Không tìm thấy phiên phân tích.')
    if session['status'] != 'completed':
        raise HTTPException(409, 'Phiên phải có trạng thái completed để xuất báo cáo.')
    metrics = (await db.execute(text('SELECT * FROM session_metrics WHERE session_id=:session_id'), params)).mappings().first()
    details = (await db.execute(text('''
        SELECT f.frame_index,f.timestamp_sec,f.frame_path,d.id AS face_id,
               d.bbox_x,d.bbox_y,d.bbox_width,d.bbox_height,d.face_confidence,
               d.face_image_path,e.emotion,e.score,e.is_dominant
        FROM frame_analyses f
        LEFT JOIN detected_faces d ON d.frame_id=f.id
        LEFT JOIN emotion_scores e ON e.face_id=d.id
        WHERE f.session_id=:session_id
        ORDER BY f.frame_index,d.id,e.emotion,e.id
    '''), params)).mappings().all()
    content = await run_in_threadpool(build_excel, dict(session), dict(metrics) if metrics else None, [dict(r) for r in details])
    # Đường dẫn tuyệt đối để tải được dù backend chạy từ thư mục khác.
    root = Path(os.environ.get('REPORT_OUTPUT_DIR', 'outputs/reports')).resolve()
    path = root / f'emotion_report_{session_id}_{uuid4().hex}.xlsx'
    await run_in_threadpool(save_file, path, content)
    try:
        await db.execute(text('''INSERT INTO reports(session_id,report_type,file_path)
            VALUES (:session_id,'excel',:file_path)'''),
            {'session_id': session_id, 'file_path': str(path)})
        await db.commit()
    except Exception:
        await db.rollback()
        await run_in_threadpool(path.unlink, missing_ok=True)
        raise
    return path
